import io
import os
import tempfile
from pathlib import Path

TEST_DB = Path(tempfile.gettempdir()) / f"knowflow-test-{os.getpid()}.db"
TEST_DB.unlink(missing_ok=True)
os.environ.pop("DATABASE_URL", None)
os.environ["SQLITE_PATH"] = str(TEST_DB)
os.environ["DEMO_AI_MODE"] = "1"

from fastapi.testclient import TestClient
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.main import app
import app.main as main


client = TestClient(app)


def test_health_and_demo_login():
    assert client.get("/api/v1/health").json()["status"] == "ok"
    response = client.post("/api/v1/auth/demo")
    assert response.status_code == 200
    assert response.json()["access_token"]


def test_knowledge_chat_plan_quiz_flow():
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    kb = client.post("/api/v1/knowledge-bases", json={"name": "Python Basics"}, headers=headers)
    assert kb.status_code == 201
    kb_id = kb.json()["id"]

    upload = client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("notes.md", b"# Python\nPython is a programming language.\nUse lists for ordered values.", "text/markdown")},
        headers=headers,
    )
    assert upload.status_code == 201
    assert upload.json()["status"] == "indexed"

    answer = client.post(
        "/api/v1/chat",
        json={"knowledge_base_id": kb_id, "message": "What is Python?"},
        headers=headers,
    )
    assert answer.status_code == 200
    assert answer.json()["citations"]
    assert "Python" in answer.json()["content"]

    plan = client.post("/api/v1/learning/plans", json={"goal": "Learn Python", "days": 7}, headers=headers)
    assert plan.status_code == 201
    assert len(plan.json()["tasks"]) == 7

    quiz = client.post(f"/api/v1/quizzes?knowledge_base_id={kb_id}", headers=headers)
    assert quiz.status_code == 201
    question = quiz.json()["questions"][0]
    assert "answer" not in question
    result = client.post(
        f"/api/v1/quizzes/{quiz.json()['id']}/submit",
        json={"answers": {question["id"]: question["options"][0]}},
        headers=headers,
    )
    assert result.status_code == 200
    assert result.json()["mastery_score"] > 0
    stats = client.get("/api/v1/stats", headers=headers).json()
    assert "learning_minutes" not in stats
    assert "streak_days" not in stats


def test_permission_boundary():
    assert client.get("/api/v1/users/me").status_code == 401


def test_repository_and_sse_security_boundaries():
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    assert client.post("/api/v1/repositories/import", json={"url": "http://127.0.0.1/repo"}, headers=headers).status_code == 400
    stream = client.post("/api/v1/chat/stream", json={"message": "RAG"}, headers=headers)
    assert stream.status_code == 200
    assert "event: done" in stream.text


def test_conversation_id_is_owned_by_requesting_user():
    user_a = client.post(
        "/api/v1/auth/register",
        json={"email": "owner@example.com", "password": "password1", "name": "Owner"},
    ).json()
    user_b = client.post(
        "/api/v1/auth/register",
        json={"email": "reader@example.com", "password": "password1", "name": "Reader"},
    ).json()
    headers_a = {"Authorization": f"Bearer {user_a['access_token']}"}
    headers_b = {"Authorization": f"Bearer {user_b['access_token']}"}
    created = client.post("/api/v1/chat", json={"message": "private"}, headers=headers_a)
    conversation_id = created.json()["conversation_id"]

    response = client.post(
        "/api/v1/chat",
        json={"message": "write into another account", "conversation_id": conversation_id},
        headers=headers_b,
    )
    assert response.status_code == 404


def _pdf_fixture() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=300, height=300)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {NameObject("/F1"): writer._add_object(font)}
            )
        }
    )
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 200 Td (PDF fixture text) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def _blank_pdf_fixture() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=300, height=300)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


def test_real_document_formats_are_extracted():
    doc = Document()
    doc.add_paragraph("DOCX fixture text")
    docx_buffer = io.BytesIO()
    doc.save(docx_buffer)

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text = "PPTX fixture text"
    pptx_buffer = io.BytesIO()
    presentation.save(pptx_buffer)

    token = client.post("/api/v1/auth/demo").json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    kb_id = client.post("/api/v1/knowledge-bases", json={"name": "Formats"}, headers=headers).json()["id"]
    for filename, payload, expected in [
        ("notes.docx", docx_buffer.getvalue(), "DOCX fixture text"),
        ("slides.pptx", pptx_buffer.getvalue(), "PPTX fixture text"),
        ("notes.pdf", _pdf_fixture(), "PDF fixture text"),
    ]:
        response = client.post(
            f"/api/v1/knowledge-bases/{kb_id}/documents",
            files={"file": (filename, payload)},
            headers=headers,
        )
        assert response.status_code == 201
        document = client.get(f"/api/v1/documents/{response.json()['id']}", headers=headers).json()
        assert expected in document["content"]

    pdf_response = client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("blank.pdf", _blank_pdf_fixture(), "application/pdf")},
        headers=headers,
    )
    assert pdf_response.status_code == 422

    text_response = client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("broken.md", b"\xff\xfe", "text/markdown")},
        headers=headers,
    )
    assert text_response.status_code == 422


def test_document_delete_requires_ownership():
    owner = client.post(
        "/api/v1/auth/register",
        json={"email": "document-owner@example.com", "password": "password1", "name": "Owner"},
    ).json()
    other = client.post(
        "/api/v1/auth/register",
        json={"email": "document-reader@example.com", "password": "password1", "name": "Reader"},
    ).json()
    owner_headers = {"Authorization": f"Bearer {owner['access_token']}"}
    other_headers = {"Authorization": f"Bearer {other['access_token']}"}
    kb_id = client.post(
        "/api/v1/knowledge-bases", json={"name": "Delete me"}, headers=owner_headers
    ).json()["id"]
    uploaded = client.post(
        f"/api/v1/knowledge-bases/{kb_id}/documents",
        files={"file": ("notes.md", b"delete this", "text/markdown")},
        headers=owner_headers,
    )
    document_id = uploaded.json()["id"]

    assert client.delete(f"/api/v1/documents/{document_id}", headers=other_headers).status_code == 404
    assert client.delete(f"/api/v1/documents/{document_id}", headers=owner_headers).status_code == 200
    assert client.get(f"/api/v1/documents/{document_id}", headers=owner_headers).status_code == 404


def test_repository_import_scans_static_files_and_indexes_readme(monkeypatch, tmp_path):
    repo = tmp_path / "demo-repo"
    repo.mkdir()
    (repo / "README.md").write_text("# Demo Repo\nUses FastAPI and SQLite.", encoding="utf-8")
    (repo / "requirements.txt").write_text("fastapi==1.0\n", encoding="utf-8")
    (repo / "app.py").write_text("def main():\n    return 'ok'\n", encoding="utf-8")
    monkeypatch.setattr(main, "clone_repository", lambda url, target: repo)
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    response = client.post(
        "/api/v1/repositories/import",
        json={"url": "https://github.com/example/demo"},
        headers=headers,
    )
    assert response.status_code == 201
    payload = response.json()
    assert "FastAPI" in payload["analysis"]["learning_summary"]
    assert payload["knowledge_base_id"]
    documents = client.get(
        f"/api/v1/knowledge-bases/{payload['knowledge_base_id']}/documents",
        headers=headers,
    ).json()
    assert any(document["filename"] == "README.md" for document in documents)


def test_memory_is_owned_searchable_and_deletable():
    owner = client.post(
        "/api/v1/auth/register",
        json={"email": "memory-owner@example.com", "password": "password1", "name": "Owner"},
    ).json()
    other = client.post(
        "/api/v1/auth/register",
        json={"email": "memory-other@example.com", "password": "password1", "name": "Other"},
    ).json()
    owner_headers = {"Authorization": f"Bearer {owner['access_token']}"}
    other_headers = {"Authorization": f"Bearer {other['access_token']}"}
    created = client.post(
        "/api/v1/memories",
        json={"memory_type": "learning_goal", "content": "重点学习 Flutter", "importance": 0.8},
        headers=owner_headers,
    )
    assert created.status_code == 201
    memory_id = created.json()["id"]
    assert client.get("/api/v1/memories?query=Flutter", headers=owner_headers).json()[0]["id"] == memory_id
    assert client.get(f"/api/v1/memories/{memory_id}", headers=other_headers).status_code == 404
    assert client.delete(f"/api/v1/memories/{memory_id}", headers=other_headers).status_code == 404
    assert client.delete(f"/api/v1/memories/{memory_id}", headers=owner_headers).status_code == 200


def test_agent_executes_validated_tool_calls_and_records_audit():
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    response = client.post(
        "/api/v1/agents/run",
        json={
            "message": "remember this",
            "tool_calls": [
                {
                    "name": "save_memory",
                    "arguments": {"memory_type": "preference", "content": "喜欢短练习"},
                }
            ],
        },
        headers=headers,
    )
    assert response.status_code == 200
    assert response.json()["tool_calls"][0]["status"] == "completed"
    assert client.get("/api/v1/agents/runs", headers=headers).json()[0]["tool_name"] == "save_memory"


def test_repository_rejects_public_host_resolving_to_private_ip(monkeypatch):
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    monkeypatch.setattr(main.socket, "getaddrinfo", lambda *args, **kwargs: [(None, None, None, None, ("10.0.0.5", 443))])
    response = client.post(
        "/api/v1/repositories/import",
        json={"url": "https://github.com/example/private"},
        headers=headers,
    )
    assert response.status_code == 400


def test_agent_requires_provider_or_structured_tool_calls():
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    response = client.post(
        "/api/v1/agents/run",
        json={"message": "What did I learn?"},
        headers={"Authorization": f"Bearer {token}"},
    )
    assert response.status_code == 503


def test_agent_tool_schema_exposes_required_arguments():
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    tools = client.get(
        "/api/v1/agents/tools",
        headers={"Authorization": f"Bearer {token}"},
    ).json()
    search = next(tool for tool in tools if tool["name"] == "search_knowledge")
    assert search["json_schema"]["required"] == ["query"]


def test_configured_provider_can_call_a_tool_then_return_final_text(monkeypatch):
    token = client.post("/api/v1/auth/demo").json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}
    responses = iter([
        {"choices": [{"message": {"role": "assistant", "tool_calls": [{"id": "call-1", "function": {"name": "search_memory", "arguments": '{"query":"Flutter"}'}}]}}]},
        {"choices": [{"message": {"role": "assistant", "content": "你最近重点学习 Flutter。"}}]},
    ])
    monkeypatch.setattr(main, "provider_configured", lambda: True)
    monkeypatch.setattr(main, "provider_completion", lambda messages: next(responses))
    response = client.post("/api/v1/agents/run", json={"message": "我最近学什么？"}, headers=headers)
    assert response.status_code == 200
    assert response.json()["content"] == "你最近重点学习 Flutter。"
    assert response.json()["tool_calls"][0]["name"] == "search_memory"
