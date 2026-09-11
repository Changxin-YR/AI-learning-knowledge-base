from __future__ import annotations

import ipaddress
import io
import json
import hashlib
import math
import os
import re
import secrets
import sqlite3
import socket
import subprocess
import tempfile
import uuid
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from urllib.request import Request, urlopen

import jwt
from argon2 import PasswordHasher
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)
DB_PATH = Path(os.getenv("SQLITE_PATH", DATA_DIR / "knowflow.db"))
DEFAULT_SECRET = "knowflow-local-secret-change-me-32-bytes"
SECRET = os.getenv("JWT_SECRET", DEFAULT_SECRET)
if os.getenv("APP_ENV", "development").lower() in {"production", "prod"} and SECRET == DEFAULT_SECRET:
    raise RuntimeError("JWT_SECRET must be changed in production")
TOKEN_TTL_MINUTES = int(os.getenv("ACCESS_TOKEN_MINUTES", "60"))
PH = PasswordHasher()
RAG_MODE = os.getenv("RAG_MODE", "hybrid").lower()
QDRANT_URL = os.getenv("QDRANT_URL", "http://127.0.0.1:6333").rstrip("/")
QDRANT_COLLECTION = os.getenv("QDRANT_COLLECTION", "knowflow_chunks")
EMBEDDING_DIMENSIONS = int(os.getenv("EMBEDDING_DIMENSIONS", "256"))
VECTOR_SCORE_THRESHOLD = float(os.getenv("VECTOR_SCORE_THRESHOLD", "0.35"))
HYBRID_STRONG_VECTOR_THRESHOLD = float(os.getenv("HYBRID_STRONG_VECTOR_THRESHOLD", "0.415"))
HYBRID_AGREEMENT_TOP_K = int(os.getenv("HYBRID_AGREEMENT_TOP_K", "2"))


def now() -> str:
    return datetime.now(timezone.utc).isoformat()


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB_PATH, check_same_thread=False)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys=ON")
    return conn


def migrate() -> None:
    schema = """
    CREATE TABLE IF NOT EXISTS schema_version(version INTEGER PRIMARY KEY);
    CREATE TABLE IF NOT EXISTS users(id TEXT PRIMARY KEY, email TEXT UNIQUE NOT NULL, name TEXT NOT NULL, password_hash TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS refresh_tokens(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), token_hash TEXT UNIQUE NOT NULL, revoked INTEGER NOT NULL DEFAULT 0, expires_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS knowledge_bases(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), name TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, knowledge_base_id TEXT NOT NULL REFERENCES knowledge_bases(id) ON DELETE CASCADE, filename TEXT NOT NULL, mime_type TEXT NOT NULL, content TEXT NOT NULL, status TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS chunks(id TEXT PRIMARY KEY, document_id TEXT NOT NULL REFERENCES documents(id) ON DELETE CASCADE, content TEXT NOT NULL, position INTEGER NOT NULL, metadata TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS conversations(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), title TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS messages(id TEXT PRIMARY KEY, conversation_id TEXT NOT NULL REFERENCES conversations(id) ON DELETE CASCADE, role TEXT NOT NULL, content TEXT NOT NULL, citations TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS study_plans(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), goal TEXT NOT NULL, days INTEGER NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS study_tasks(id TEXT PRIMARY KEY, plan_id TEXT NOT NULL REFERENCES study_plans(id) ON DELETE CASCADE, title TEXT NOT NULL, due_date TEXT NOT NULL, completed INTEGER NOT NULL DEFAULT 0);
    CREATE TABLE IF NOT EXISTS quizzes(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), knowledge_base_id TEXT, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS quiz_questions(id TEXT PRIMARY KEY, quiz_id TEXT NOT NULL REFERENCES quizzes(id) ON DELETE CASCADE, prompt TEXT NOT NULL, options TEXT NOT NULL, answer TEXT NOT NULL, knowledge TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS quiz_attempts(id TEXT PRIMARY KEY, quiz_id TEXT NOT NULL REFERENCES quizzes(id), user_id TEXT NOT NULL REFERENCES users(id), score REAL NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS mastery(user_id TEXT NOT NULL REFERENCES users(id), knowledge TEXT NOT NULL, score REAL NOT NULL, updated_at TEXT NOT NULL, PRIMARY KEY(user_id, knowledge));
    CREATE TABLE IF NOT EXISTS repository_imports(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id), url TEXT NOT NULL, status TEXT NOT NULL, analysis TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS memories(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, memory_type TEXT NOT NULL, content TEXT NOT NULL, importance REAL NOT NULL, source TEXT NOT NULL, metadata TEXT NOT NULL, created_at TEXT NOT NULL, updated_at TEXT NOT NULL, last_accessed_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS agent_runs(id TEXT PRIMARY KEY, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, conversation_id TEXT, message TEXT NOT NULL, status TEXT NOT NULL, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS tool_calls(id TEXT PRIMARY KEY, run_id TEXT NOT NULL REFERENCES agent_runs(id) ON DELETE CASCADE, user_id TEXT NOT NULL REFERENCES users(id) ON DELETE CASCADE, tool_name TEXT NOT NULL, arguments TEXT NOT NULL, result_status TEXT NOT NULL, result TEXT NOT NULL, created_at TEXT NOT NULL);
    """
    with db() as conn:
        conn.executescript(schema)
        conn.execute("INSERT OR IGNORE INTO schema_version(version) VALUES (1)")


migrate()
app = FastAPI(title="KnowFlow AI API", version="1.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=os.getenv("CORS_ORIGINS", "*").split(","),
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def request_id(request, call_next):
    trace_id = request.headers.get("X-Request-ID", str(uuid.uuid4()))
    response = await call_next(request)
    response.headers["X-Request-ID"] = trace_id
    return response


class KBIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)


class ChatIn(BaseModel):
    message: str = Field(min_length=1, max_length=4000)
    knowledge_base_id: str | None = None
    conversation_id: str | None = None


class PlanIn(BaseModel):
    goal: str = Field(min_length=1, max_length=200)
    days: int = Field(default=7, ge=1, le=14)


class SubmitIn(BaseModel):
    answers: dict[str, str]


class MemoryIn(BaseModel):
    memory_type: str = Field(default="user_note", pattern="^(preference|learning_goal|weak_knowledge|project_context|user_note)$")
    content: str = Field(min_length=1, max_length=4000)
    importance: float = Field(default=0.5, ge=0, le=1)
    source: str = Field(default="user", max_length=80)
    metadata: dict[str, Any] = Field(default_factory=dict)


class RepositoryIn(BaseModel):
    url: str = Field(min_length=1, max_length=500)
    knowledge_base_id: str | None = None


class AgentCallIn(BaseModel):
    name: str = Field(min_length=1, max_length=80)
    arguments: dict[str, Any] = Field(default_factory=dict)


class AgentRunIn(BaseModel):
    message: str = Field(min_length=1, max_length=4000)
    conversation_id: str | None = None
    tool_calls: list[AgentCallIn] = Field(default_factory=list, max_length=8)


def token_for(user_id: str) -> str:
    return jwt.encode({"sub": user_id, "exp": datetime.now(timezone.utc) + timedelta(minutes=TOKEN_TTL_MINUTES)}, SECRET, algorithm="HS256")


def current_user(authorization: str | None = Header(default=None)) -> sqlite3.Row:
    if not authorization or not authorization.lower().startswith("bearer "):
        raise HTTPException(401, "Authentication required")
    try:
        payload = jwt.decode(authorization.split(" ", 1)[1], SECRET, algorithms=["HS256"])
        user_id = payload["sub"]
    except (jwt.PyJWTError, KeyError):
        raise HTTPException(401, "Invalid or expired token")
    with db() as conn:
        user = conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if not user:
        raise HTTPException(401, "User not found")
    return user


def owned_kb(user_id: str, kb_id: str) -> sqlite3.Row:
    with db() as conn:
        row = conn.execute("SELECT * FROM knowledge_bases WHERE id=? AND user_id=?", (kb_id, user_id)).fetchone()
    if not row:
        raise HTTPException(404, "Knowledge base not found")
    return row


def owned_conversation(user_id: str, conversation_id: str) -> sqlite3.Row:
    with db() as conn:
        row = conn.execute(
            "SELECT * FROM conversations WHERE id=? AND user_id=?",
            (conversation_id, user_id),
        ).fetchone()
    if not row:
        raise HTTPException(404, "Conversation not found")
    return row


def safe_filename(name: str) -> str:
    clean = Path(name).name
    if not clean or clean in {".", ".."} or len(clean) > 180:
        raise HTTPException(400, "Invalid filename")
    return clean


def parse_content(filename: str, payload: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx", ".pptx", ".md", ".txt", ".py", ".dart", ".js", ".ts", ".json", ".yaml", ".yml"}:
        raise HTTPException(415, "Unsupported file type")
    if len(payload) > 10 * 1024 * 1024:
        raise HTTPException(413, "File too large")
    try:
        if suffix == ".pdf":
            from pypdf import PdfReader

            text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(payload)).pages)
        elif suffix == ".docx":
            from docx import Document

            document = Document(io.BytesIO(payload))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            text = "\n".join(parts)
        elif suffix == ".pptx":
            from pptx import Presentation

            presentation = Presentation(io.BytesIO(payload))
            text = "\n".join(
                shape.text
                for slide in presentation.slides
                for shape in slide.shapes
                if hasattr(shape, "text")
            )
        else:
            text = payload.decode("utf-8", errors="strict")
    except UnicodeDecodeError as exc:
        raise HTTPException(422, "Text files must be UTF-8") from exc
    except ImportError as exc:
        raise HTTPException(500, "Document parser dependency is not installed") from exc
    except Exception as exc:
        raise HTTPException(422, "Document could not be parsed") from exc
    if not text.strip():
        raise HTTPException(422, "Document contains no extractable text")
    return text.strip()


def make_chunks(text: str) -> list[str]:
    blocks = [x.strip() for x in re.split(r"\n\s*\n|(?=^#{1,6}\s)", text, flags=re.M) if x.strip()]
    chunks: list[str] = []
    for block in blocks:
        if len(block) <= 900:
            chunks.append(block)
        else:
            chunks.extend(block[i : i + 900] for i in range(0, len(block), 900))
    return chunks or [text[:900]]


def search_tokens(text: str) -> set[str]:
    words = re.findall(r"[a-z0-9_]+|[\u4e00-\u9fff]", text.lower())
    compact = "".join(char for char in text.lower() if "\u4e00" <= char <= "\u9fff")
    return set(words) | {compact[index : index + size] for size in (2, 3) for index in range(max(0, len(compact) - size + 1))}


class EmbeddingProvider:
    dimensions: int

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        raise NotImplementedError

    def embed_query(self, text: str) -> list[float]:
        return self.embed_documents([text])[0]


class LocalEmbeddingProvider(EmbeddingProvider):
    """Small, deterministic local vectorizer for offline/dev use."""

    def __init__(self, dimensions: int = EMBEDDING_DIMENSIONS):
        if dimensions < 8:
            raise ValueError("Embedding dimensions must be at least 8")
        self.dimensions = dimensions

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        vectors = []
        for text in texts:
            vector = [0.0] * self.dimensions
            for token in search_tokens(text):
                digest = hashlib.sha256(token.encode("utf-8")).digest()
                index = int.from_bytes(digest[:4], "big") % self.dimensions
                vector[index] += 1.0 if len(token) > 1 else 0.5
            norm = math.sqrt(sum(value * value for value in vector))
            vectors.append([value / norm for value in vector] if norm else vector)
        return vectors


class OpenAIEmbeddingProvider(EmbeddingProvider):
    def __init__(self):
        self.endpoint = os.getenv("EMBEDDING_BASE_URL", "").rstrip("/")
        self.api_key = os.getenv("EMBEDDING_API_KEY", "")
        self.model = os.getenv("EMBEDDING_MODEL", "")
        self.dimensions = int(os.getenv("EMBEDDING_DIMENSIONS", str(EMBEDDING_DIMENSIONS)))

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        endpoint = self.endpoint
        if not endpoint.endswith("/embeddings"):
            endpoint += "/embeddings"
        payload = json.dumps({"model": self.model, "input": texts}).encode("utf-8")
        request = Request(endpoint, data=payload, headers={"Authorization": f"Bearer {self.api_key}", "Content-Type": "application/json"}, method="POST")
        try:
            with urlopen(request, timeout=20) as response:
                data = json.loads(response.read().decode("utf-8"))
            vectors = [item["embedding"] for item in sorted(data["data"], key=lambda item: item.get("index", 0))]
            if not vectors or any(len(vector) != len(vectors[0]) for vector in vectors):
                raise ValueError("Embedding provider returned invalid vectors")
            self.dimensions = len(vectors[0])
            return vectors
        except Exception as exc:
            raise RuntimeError("Configured embedding provider failed") from exc


def embedding_provider() -> EmbeddingProvider:
    provider = os.getenv("EMBEDDING_PROVIDER", "local").lower()
    if provider == "openai" or (provider == "auto" and all(os.getenv(key) for key in ("EMBEDDING_BASE_URL", "EMBEDDING_API_KEY", "EMBEDDING_MODEL"))):
        return OpenAIEmbeddingProvider()
    return LocalEmbeddingProvider()


def cosine_similarity(left: list[float], right: list[float]) -> float:
    if len(left) != len(right):
        return 0.0
    denominator = math.sqrt(sum(value * value for value in left) * sum(value * value for value in right))
    return sum(a * b for a, b in zip(left, right)) / denominator if denominator else 0.0


def _qdrant_request(method: str, path: str, payload: dict[str, Any] | None = None) -> dict[str, Any]:
    data = json.dumps(payload).encode("utf-8") if payload is not None else None
    request = Request(f"{QDRANT_URL}{path}", data=data, headers={"Content-Type": "application/json"}, method=method)
    with urlopen(request, timeout=5) as response:
        body = response.read()
    return json.loads(body.decode("utf-8")) if body else {}


def qdrant_available() -> bool:
    try:
        _qdrant_request("GET", "/collections")
        return True
    except Exception:
        return False


def ensure_qdrant_collection(dimensions: int) -> None:
    try:
        existing = _qdrant_request("GET", f"/collections/{QDRANT_COLLECTION}")
        size = existing.get("result", {}).get("config", {}).get("params", {}).get("vectors", {}).get("size")
        if size and size != dimensions:
            raise RuntimeError(f"Qdrant collection dimension is {size}, expected {dimensions}")
        return
    except Exception as exc:
        if "HTTP Error 404" not in str(exc):
            raise
    _qdrant_request("PUT", f"/collections/{QDRANT_COLLECTION}", {"vectors": {"size": dimensions, "distance": "Cosine"}})


def qdrant_upsert(points: list[dict[str, Any]], dimensions: int) -> bool:
    if not points:
        return True
    try:
        ensure_qdrant_collection(dimensions)
        _qdrant_request("PUT", f"/collections/{QDRANT_COLLECTION}/points?wait=true", {"points": points})
        return True
    except Exception:
        return False


def qdrant_delete(filter_body: dict[str, Any]) -> bool:
    try:
        _qdrant_request("POST", f"/collections/{QDRANT_COLLECTION}/points/delete?wait=true", {"filter": filter_body})
        return True
    except Exception:
        return False


def qdrant_search_with_status(vector: list[float], user_id: str, kb_id: str | None, limit: int, threshold: float) -> tuple[list[tuple[str, float]], bool]:
    must = [{"key": "user_id", "match": {"value": user_id}}]
    if kb_id:
        must.append({"key": "knowledge_base_id", "match": {"value": kb_id}})
    try:
        ensure_qdrant_collection(len(vector))
        result = _qdrant_request("POST", f"/collections/{QDRANT_COLLECTION}/points/search", {"vector": vector, "limit": limit, "with_payload": True, "score_threshold": threshold, "filter": {"must": must}})
        return [(str(item["id"]), float(item.get("score", 0))) for item in result.get("result", [])], True
    except Exception:
        return [], False


def qdrant_search(vector: list[float], user_id: str, kb_id: str | None, limit: int, threshold: float) -> list[tuple[str, float]]:
    return qdrant_search_with_status(vector, user_id, kb_id, limit, threshold)[0]


def _chunk_rows(user_id: str, kb_id: str | None, ids: list[str]) -> list[sqlite3.Row]:
    if not ids:
        return []
    placeholders = ",".join("?" for _ in ids)
    params: list[Any] = [*ids, user_id]
    where = f"c.id IN ({placeholders}) AND k.user_id=?"
    if kb_id:
        owned_kb(user_id, kb_id)
        where += " AND d.knowledge_base_id=?"
        params.append(kb_id)
    with db() as conn:
        rows = conn.execute(f"SELECT c.*, d.filename, d.knowledge_base_id FROM chunks c JOIN documents d ON d.id=c.document_id JOIN knowledge_bases k ON k.id=d.knowledge_base_id WHERE {where}", params).fetchall()
    by_id = {row["id"]: row for row in rows}
    return [by_id[chunk_id] for chunk_id in ids if chunk_id in by_id]


def retrieve_lexical(user_id: str, kb_id: str | None, query: str, limit: int = 4) -> list[sqlite3.Row]:
    terms = search_tokens(query)
    with db() as conn:
        if kb_id:
            owned_kb(user_id, kb_id)
            rows = conn.execute("SELECT c.*, d.filename, d.knowledge_base_id FROM chunks c JOIN documents d ON d.id=c.document_id WHERE d.knowledge_base_id=?", (kb_id,)).fetchall()
        else:
            rows = conn.execute("SELECT c.*, d.filename, d.knowledge_base_id FROM chunks c JOIN documents d ON d.id=c.document_id JOIN knowledge_bases k ON k.id=d.knowledge_base_id WHERE k.user_id=?", (user_id,)).fetchall()
    scored = [(len(terms & search_tokens(row["content"])), row) for row in rows]
    return [row for score, row in sorted(scored, key=lambda item: item[0], reverse=True) if score > 0][:limit]


def retrieve_vector(user_id: str, kb_id: str | None, query: str, limit: int = 4) -> list[sqlite3.Row]:
    provider = embedding_provider()
    matches = qdrant_search(provider.embed_query(query), user_id, kb_id, limit, float(os.getenv("VECTOR_SCORE_THRESHOLD", str(VECTOR_SCORE_THRESHOLD))))
    return _chunk_rows(user_id, kb_id, [chunk_id for chunk_id, _score in matches])


def retrieve_vector_evidence(user_id: str, kb_id: str | None, query: str, limit: int = 4) -> tuple[list[tuple[str, float]], bool]:
    try:
        provider = embedding_provider()
        vector = provider.embed_query(query)
    except Exception:
        return [], False
    threshold = float(os.getenv("VECTOR_SCORE_THRESHOLD", str(VECTOR_SCORE_THRESHOLD)))
    return qdrant_search_with_status(vector, user_id, kb_id, limit, threshold)


def rrf_fusion(rankings: list[list[str]], k: int = 60, limit: int = 4) -> list[tuple[str, float]]:
    scores: dict[str, float] = {}
    for ranking in rankings:
        for rank, item_id in enumerate(ranking, 1):
            scores[item_id] = scores.get(item_id, 0.0) + 1 / (k + rank)
    return sorted(scores.items(), key=lambda item: (-item[1], item[0]))[:limit]


def hybrid_confidence_gate_enabled() -> bool:
    configured = os.getenv("HYBRID_CONFIDENCE_GATE", "auto").strip().lower()
    if configured in {"1", "true", "yes", "on"}:
        return True
    if configured in {"0", "false", "no", "off"}:
        return False
    return os.getenv("EMBEDDING_PROVIDER", "local").lower() != "local"


def hybrid_confidence_allows(
    lexical_ids: list[str],
    vector_matches: list[tuple[str, float]],
    strong_threshold: float | None = None,
    agreement_top_k: int | None = None,
) -> bool:
    if not vector_matches:
        return False
    threshold = strong_threshold if strong_threshold is not None else float(os.getenv("HYBRID_STRONG_VECTOR_THRESHOLD", str(HYBRID_STRONG_VECTOR_THRESHOLD)))
    if max(score for _item_id, score in vector_matches) >= threshold:
        return True
    top_k = max(1, agreement_top_k if agreement_top_k is not None else int(os.getenv("HYBRID_AGREEMENT_TOP_K", str(HYBRID_AGREEMENT_TOP_K))))
    lexical_top = set(lexical_ids[:top_k])
    vector_top = {item_id for item_id, _score in vector_matches[:top_k]}
    return bool(lexical_top & vector_top)


def semantic_backed_rrf_ids(
    lexical_ids: list[str],
    vector_matches: list[tuple[str, float]],
    limit: int = 4,
    strong_threshold: float | None = None,
    agreement_top_k: int | None = None,
) -> list[str]:
    if not hybrid_confidence_allows(lexical_ids, vector_matches, strong_threshold, agreement_top_k):
        return []
    vector_ids = [item_id for item_id, _score in vector_matches]
    vector_set = set(vector_ids)
    candidate_limit = max(limit * 2, len(set(lexical_ids) | vector_set))
    fused = rrf_fusion([lexical_ids, vector_ids], limit=candidate_limit)
    return [item_id for item_id, _score in fused if item_id in vector_set][:limit]


def retrieve(user_id: str, kb_id: str | None, query: str, limit: int = 4) -> list[sqlite3.Row]:
    mode = os.getenv("RAG_MODE", RAG_MODE).lower()
    if mode == "lexical":
        return retrieve_lexical(user_id, kb_id, query, limit)
    if mode == "vector":
        return retrieve_vector(user_id, kb_id, query, limit)
    lexical = retrieve_lexical(user_id, kb_id, query, limit)
    vector_matches, vector_available = retrieve_vector_evidence(user_id, kb_id, query, limit)
    if not vector_available:
        return lexical
    candidate_rows = _chunk_rows(user_id, kb_id, [item_id for item_id, _score in vector_matches])
    owned_ids = {row["id"] for row in candidate_rows}
    vector_matches = [(item_id, score) for item_id, score in vector_matches if item_id in owned_ids]
    lexical_ids = [row["id"] for row in lexical]
    if hybrid_confidence_gate_enabled():
        fused_ids = semantic_backed_rrf_ids(lexical_ids, vector_matches, limit=limit)
    else:
        vector_ids = [row["id"] for row in candidate_rows]
        fused_ids = [item_id for item_id, _score in rrf_fusion([lexical_ids, vector_ids], limit=limit)]
    return _chunk_rows(user_id, kb_id, fused_ids)


def index_chunks(user_id: str, kb_id: str, document_id: str, chunks: list[tuple[str, str, int, str]] | list[tuple[str, str, int]]) -> bool:
    if not chunks:
        return True
    try:
        provider = embedding_provider()
        vectors = provider.embed_documents([chunk for _chunk_id, chunk, _position, *_ in chunks])
        points = []
        for chunk_record, vector in zip(chunks, vectors):
            chunk_id, _chunk, _position, *chunk_document = chunk_record
            points.append({
                "id": chunk_id,
                "vector": vector,
                "payload": {"user_id": user_id, "knowledge_base_id": kb_id, "document_id": chunk_document[0] if chunk_document else document_id, "chunk_id": chunk_id},
            })
        return qdrant_upsert(points, provider.dimensions)
    except Exception:
        return False


PUBLIC_REPOSITORY_HOSTS = {"github.com", "gitlab.com", "gitee.com"}
TEXT_FILE_SUFFIXES = {".md", ".txt", ".py", ".dart", ".js", ".ts", ".json", ".yaml", ".yml", ".toml", ".xml", ".java", ".kt", ".go", ".rs", ".rb", ".php", ".cs", ".c", ".h", ".cpp", ".gradle", ".properties"}
DEPENDENCY_FILES = {"package.json", "pubspec.yaml", "requirements.txt", "pyproject.toml", "pom.xml", "build.gradle", "build.gradle.kts", "Cargo.toml", "go.mod"}


def public_repository_url(raw_url: str):
    parsed = urlparse(raw_url.strip())
    host = (parsed.hostname or "").lower().rstrip(".")
    if parsed.scheme not in {"http", "https"} or host not in PUBLIC_REPOSITORY_HOSTS or parsed.username or parsed.password:
        raise HTTPException(400, "Only public GitHub, GitLab, or Gitee repositories are allowed")
    try:
        addresses = {item[4][0] for item in socket.getaddrinfo(host, parsed.port or 443, type=socket.SOCK_STREAM)}
    except OSError as exc:
        raise HTTPException(400, "Repository host could not be resolved") from exc
    for address in addresses:
        ip = ipaddress.ip_address(address)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_multicast or ip.is_reserved or ip.is_unspecified:
            raise HTTPException(400, "Repository host resolves to a non-public address")
    return parsed


def clone_repository(url: str, target: Path) -> Path:
    public_repository_url(url)
    timeout = int(os.getenv("REPOSITORY_CLONE_TIMEOUT_SECONDS", "30"))
    try:
        subprocess.run(
            ["git", "-c", "http.followRedirects=false", "clone", "--depth", "1", "--no-recurse-submodules", url, str(target)],
            check=True,
            capture_output=True,
            text=True,
            timeout=timeout,
        )
    except (subprocess.SubprocessError, OSError) as exc:
        raise HTTPException(422, "Repository could not be cloned") from exc
    return target


def repository_files(root: Path) -> list[tuple[Path, int]]:
    files: list[tuple[Path, int]] = []
    total_size = 0
    max_files = int(os.getenv("REPOSITORY_MAX_FILES", "2000"))
    max_file_size = int(os.getenv("REPOSITORY_MAX_FILE_BYTES", str(1024 * 1024)))
    max_total_size = int(os.getenv("REPOSITORY_MAX_TOTAL_BYTES", str(50 * 1024 * 1024)))
    for path in root.rglob("*"):
        if ".git" in path.parts or not path.is_file() or path.is_symlink():
            continue
        size = path.stat().st_size
        if size > max_file_size:
            continue
        total_size += size
        if total_size > max_total_size or len(files) >= max_files:
            break
        files.append((path, size))
    return files


def analyze_repository(root: Path, project_name: str) -> tuple[dict[str, Any], list[tuple[str, str]]]:
    files = repository_files(root)
    language_counts: dict[str, int] = {}
    tree: list[str] = []
    source_documents: list[tuple[str, str]] = []
    readme = ""
    dependencies: list[str] = []
    frameworks: set[str] = set()
    entry_points: list[str] = []
    language_by_suffix = {".py": "Python", ".dart": "Dart", ".js": "JavaScript", ".ts": "TypeScript", ".java": "Java", ".kt": "Kotlin", ".go": "Go", ".rs": "Rust", ".rb": "Ruby", ".php": "PHP", ".cs": "C#", ".cpp": "C++", ".c": "C"}
    for path, _ in files:
        relative = path.relative_to(root).as_posix()
        tree.append(relative)
        suffix = path.suffix.lower()
        if suffix in language_by_suffix:
            language_counts[language_by_suffix[suffix]] = language_counts.get(language_by_suffix[suffix], 0) + 1
        if path.name.lower().startswith("readme"):
            entry_points.append(relative)
            readme = path.read_text(encoding="utf-8", errors="replace")[:12000]
            source_documents.append((relative, readme))
        elif path.name in DEPENDENCY_FILES:
            content = path.read_text(encoding="utf-8", errors="replace")[:12000]
            dependencies.extend(line.strip() for line in content.splitlines() if line.strip() and not line.lstrip().startswith(("#", "//")))
            source_documents.append((relative, content))
            if path.name in {"pubspec.yaml", "build.gradle", "build.gradle.kts"}:
                frameworks.add("Flutter/Gradle")
            if path.name in {"requirements.txt", "pyproject.toml"}:
                frameworks.add("Python")
        elif suffix in TEXT_FILE_SUFFIXES and len(source_documents) < 25:
            content = path.read_text(encoding="utf-8", errors="replace")[:12000]
            source_documents.append((relative, content))
        if path.name.lower() in {"main.py", "main.dart", "index.js", "index.ts", "main.go", "main.rs"}:
            entry_points.append(relative)
    summary = (readme.strip() or f"{project_name} contains {len(tree)} static files.")[:2000]
    analysis = {
        "name": project_name,
        "languages": sorted(language_counts),
        "frameworks": sorted(frameworks),
        "dependencies": dependencies[:100],
        "entry_points": sorted(set(entry_points))[:50],
        "directory_tree": tree[:500],
        "learning_summary": summary,
    }
    return analysis, source_documents


def create_memory(user_id: str, body: MemoryIn) -> dict[str, Any]:
    memory_id = str(uuid.uuid4())
    timestamp = now()
    with db() as conn:
        conn.execute(
            "INSERT INTO memories VALUES (?,?,?,?,?,?,?,?,?,?)",
            (memory_id, user_id, body.memory_type, body.content.strip(), body.importance, body.source, json.dumps(body.metadata, ensure_ascii=False), timestamp, timestamp, timestamp),
        )
    return {"id": memory_id, "user_id": user_id, "memory_type": body.memory_type, "content": body.content.strip(), "importance": body.importance, "source": body.source, "metadata": body.metadata, "created_at": timestamp, "updated_at": timestamp, "last_accessed_at": timestamp}


def search_memories(user_id: str, query: str | None = None, limit: int = 10) -> list[dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT * FROM memories WHERE user_id=? ORDER BY importance DESC, updated_at DESC LIMIT 100", (user_id,)).fetchall()
        if query:
            terms = search_tokens(query)
            rows = [row for row in rows if terms & search_tokens(row["content"])]
        results = [dict(row) for row in rows[:limit]]
        for row in results:
            row["metadata"] = json.loads(row["metadata"])
        if results:
            conn.executemany("UPDATE memories SET last_accessed_at=? WHERE id=? AND user_id=?", [(now(), row["id"], user_id) for row in results])
    return results


@app.get("/api/v1/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "knowflow-api"}


@app.get("/api/v1/ready")
def ready() -> dict[str, Any]:
    with db() as conn:
        conn.execute("SELECT 1")
    return {"status": "ready", "database": "sqlite", "qdrant": qdrant_available(), "rag_mode": os.getenv("RAG_MODE", RAG_MODE), "embedding_provider": os.getenv("EMBEDDING_PROVIDER", "local"), "hybrid_confidence_gate": hybrid_confidence_gate_enabled(), "demo_ai": os.getenv("DEMO_AI_MODE", "1") == "1"}


@app.post("/api/v1/auth/demo")
def demo_login() -> dict[str, Any]:
    with db() as conn:
        user = conn.execute("SELECT * FROM users WHERE email=?", ("demo@knowflow.local",)).fetchone()
        if not user:
            user_id = str(uuid.uuid4())
            conn.execute("INSERT INTO users VALUES (?,?,?,?,?)", (user_id, "demo@knowflow.local", "Demo 学习者", PH.hash("KnowFlowDemo123!"), now()))
            conn.commit()
            user = conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    return {"access_token": token_for(user["id"]), "refresh_token": secrets.token_urlsafe(32), "token_type": "bearer", "user": {"id": user["id"], "email": user["email"], "name": user["name"]}}


@app.post("/api/v1/auth/register")
def register(body: dict[str, str]) -> dict[str, Any]:
    email, password, name = body.get("email", "").strip().lower(), body.get("password", ""), body.get("name", "学习者").strip()
    if "@" not in email or len(password) < 8:
        raise HTTPException(400, "Valid email and 8+ character password required")
    user_id = str(uuid.uuid4())
    try:
        with db() as conn:
            conn.execute("INSERT INTO users VALUES (?,?,?,?,?)", (user_id, email, name or "学习者", PH.hash(password), now()))
    except sqlite3.IntegrityError:
        raise HTTPException(409, "Email already registered")
    return {"access_token": token_for(user_id), "token_type": "bearer", "user": {"id": user_id, "email": email, "name": name or "学习者"}}


@app.get("/api/v1/users/me")
def me(user=Depends(current_user)) -> dict[str, Any]:
    return {"id": user["id"], "email": user["email"], "name": user["name"]}


@app.get("/api/v1/knowledge-bases")
def list_kbs(user=Depends(current_user)) -> list[dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT k.*, COUNT(d.id) AS document_count FROM knowledge_bases k LEFT JOIN documents d ON d.knowledge_base_id=k.id WHERE k.user_id=? GROUP BY k.id ORDER BY k.created_at DESC", (user["id"],)).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/v1/knowledge-bases", status_code=201)
def create_kb(body: KBIn, user=Depends(current_user)) -> dict[str, Any]:
    kb_id = str(uuid.uuid4())
    with db() as conn:
        conn.execute("INSERT INTO knowledge_bases VALUES (?,?,?,?)", (kb_id, user["id"], body.name.strip(), now()))
    return {"id": kb_id, "user_id": user["id"], "name": body.name.strip(), "document_count": 0}


@app.delete("/api/v1/knowledge-bases/{kb_id}")
def delete_kb(kb_id: str, user=Depends(current_user)) -> dict[str, bool]:
    owned_kb(user["id"], kb_id)
    qdrant_delete({"must": [{"key": "user_id", "match": {"value": user["id"]}}, {"key": "knowledge_base_id", "match": {"value": kb_id}}]})
    with db() as conn:
        conn.execute("DELETE FROM knowledge_bases WHERE id=?", (kb_id,))
    return {"deleted": True}


@app.get("/api/v1/knowledge-bases/{kb_id}/documents")
def list_documents(kb_id: str, user=Depends(current_user)) -> list[dict[str, Any]]:
    owned_kb(user["id"], kb_id)
    with db() as conn:
        rows = conn.execute("SELECT id,filename,mime_type,status,created_at FROM documents WHERE knowledge_base_id=? ORDER BY created_at DESC", (kb_id,)).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/v1/knowledge-bases/{kb_id}/documents", status_code=201)
async def upload_document(kb_id: str, file: UploadFile = File(...), user=Depends(current_user)) -> dict[str, Any]:
    owned_kb(user["id"], kb_id)
    filename = safe_filename(file.filename or "document.txt")
    payload = await file.read()
    content = parse_content(filename, payload)
    document_id = str(uuid.uuid4())
    chunks = make_chunks(content)
    chunk_records: list[tuple[str, str, int]] = []
    with db() as conn:
        conn.execute("INSERT INTO documents VALUES (?,?,?,?,?,?,?)", (document_id, kb_id, filename, file.content_type or "text/plain", content, "indexed", now()))
        for position, chunk in enumerate(chunks):
            chunk_id = str(uuid.uuid4())
            conn.execute("INSERT INTO chunks VALUES (?,?,?,?,?)", (chunk_id, document_id, chunk, position, json.dumps({"filename": filename, "section": position + 1})))
            chunk_records.append((chunk_id, chunk, position))
    vector_indexed = index_chunks(user["id"], kb_id, document_id, chunk_records)
    return {"id": document_id, "filename": filename, "status": "indexed", "chunk_count": len(chunks), "vector_indexed": vector_indexed}


@app.get("/api/v1/documents/{document_id}")
def document_detail(document_id: str, user=Depends(current_user)) -> dict[str, Any]:
    with db() as conn:
        row = conn.execute("SELECT d.* FROM documents d JOIN knowledge_bases k ON k.id=d.knowledge_base_id WHERE d.id=? AND k.user_id=?", (document_id, user["id"])).fetchone()
        chunks = conn.execute("SELECT id,content,position,metadata FROM chunks WHERE document_id=? ORDER BY position", (document_id,)).fetchall()
    if not row:
        raise HTTPException(404, "Document not found")
    return {**dict(row), "chunks": [dict(chunk) for chunk in chunks]}


@app.delete("/api/v1/documents/{document_id}")
def delete_document(document_id: str, user=Depends(current_user)) -> dict[str, bool]:
    with db() as conn:
        row = conn.execute(
            "SELECT d.id, d.knowledge_base_id FROM documents d JOIN knowledge_bases k ON k.id=d.knowledge_base_id WHERE d.id=? AND k.user_id=?",
            (document_id, user["id"]),
        ).fetchone()
        if not row:
            raise HTTPException(404, "Document not found")
        conn.execute("DELETE FROM documents WHERE id=?", (document_id,))
    qdrant_delete({"must": [{"key": "user_id", "match": {"value": user["id"]}}, {"key": "knowledge_base_id", "match": {"value": row["knowledge_base_id"]}}, {"key": "document_id", "match": {"value": document_id}}]})
    return {"deleted": True}


def citation(row: sqlite3.Row) -> dict[str, Any]:
    return {"chunk_id": row["id"], "document_id": row["document_id"], "filename": row["filename"], "section": row["position"] + 1, "preview": row["content"][:180]}


@app.post("/api/v1/chat")
def chat(body: ChatIn, user=Depends(current_user)) -> dict[str, Any]:
    if body.conversation_id:
        owned_conversation(user["id"], body.conversation_id)
    rows = retrieve(user["id"], body.knowledge_base_id, body.message)
    answer: str | None = None
    if rows and chat_generation_enabled():
        try:
            answer = provider_chat_answer(body.message, rows)
        except InsufficientEvidenceError:
            # Fail closed: the model rejected the retrieved evidence, so its citations must go with it.
            rows = []
        except HTTPException:
            # A provider outage must not remove an already-grounded citation; use the excerpt fallback below.
            answer = None
    if rows:
        answer = answer or f"{CHAT_DETERMINISTIC_PREFIX}{rows[0]['content'][:500]}"
        citations = [citation(row) for row in rows]
    else:
        answer = CHAT_NO_EVIDENCE_ANSWER
        citations = []
    conversation_id = body.conversation_id or str(uuid.uuid4())
    with db() as conn:
        if not body.conversation_id:
            conn.execute("INSERT INTO conversations VALUES (?,?,?,?)", (conversation_id, user["id"], body.message[:40], now()))
        conn.execute("INSERT INTO messages VALUES (?,?,?,?,?,?)", (str(uuid.uuid4()), conversation_id, "user", body.message, "[]", now()))
        conn.execute("INSERT INTO messages VALUES (?,?,?,?,?,?)", (str(uuid.uuid4()), conversation_id, "assistant", answer, json.dumps(citations, ensure_ascii=False), now()))
    return {"conversation_id": conversation_id, "content": answer, "citations": citations, "tool_events": [{"tool": "search_knowledge", "status": "completed", "count": len(rows)}]}


@app.post("/api/v1/chat/stream")
def chat_stream(body: ChatIn, user=Depends(current_user)) -> StreamingResponse:
    result = chat(body, user)
    def events():
        yield f"event: tool\\ndata: {json.dumps(result['tool_events'][0], ensure_ascii=False)}\\n\\n"
        yield f"event: message\\ndata: {json.dumps(result, ensure_ascii=False)}\\n\\n"
        yield "event: done\\ndata: {}\\n\\n"
    return StreamingResponse(events(), media_type="text/event-stream")


@app.get("/api/v1/conversations")
def conversations(user=Depends(current_user)) -> list[dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT * FROM conversations WHERE user_id=? ORDER BY created_at DESC", (user["id"],)).fetchall()
    return [dict(row) for row in rows]


@app.post("/api/v1/learning/plans", status_code=201)
def create_plan(body: PlanIn, user=Depends(current_user)) -> dict[str, Any]:
    plan_id = str(uuid.uuid4())
    created = datetime.now(timezone.utc)
    tasks = []
    with db() as conn:
        conn.execute("INSERT INTO study_plans VALUES (?,?,?,?,?)", (plan_id, user["id"], body.goal, body.days, now()))
        for index in range(body.days):
            task_id = str(uuid.uuid4())
            due = (created + timedelta(days=index)).date().isoformat()
            title = f"{body.goal} · 第 {index + 1} 天"
            conn.execute("INSERT INTO study_tasks VALUES (?,?,?,?,?)", (task_id, plan_id, title, due, 0))
            tasks.append({"id": task_id, "title": title, "due_date": due, "completed": False})
    return {"id": plan_id, "goal": body.goal, "days": body.days, "tasks": tasks}


@app.get("/api/v1/learning/tasks")
def list_tasks(user=Depends(current_user)) -> list[dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT t.* FROM study_tasks t JOIN study_plans p ON p.id=t.plan_id WHERE p.user_id=? ORDER BY t.due_date", (user["id"],)).fetchall()
    return [{**dict(row), "completed": bool(row["completed"])} for row in rows]


@app.patch("/api/v1/learning/tasks/{task_id}")
def update_task(task_id: str, body: dict[str, Any], user=Depends(current_user)) -> dict[str, Any]:
    completed = bool(body.get("completed", False))
    with db() as conn:
        row = conn.execute("SELECT t.* FROM study_tasks t JOIN study_plans p ON p.id=t.plan_id WHERE t.id=? AND p.user_id=?", (task_id, user["id"])).fetchone()
        if not row:
            raise HTTPException(404, "Task not found")
        conn.execute("UPDATE study_tasks SET completed=? WHERE id=?", (int(completed), task_id))
    return {"id": task_id, "completed": completed}


@app.post("/api/v1/quizzes", status_code=201)
def create_quiz(knowledge_base_id: str | None = None, user=Depends(current_user)) -> dict[str, Any]:
    if knowledge_base_id:
        owned_kb(user["id"], knowledge_base_id)
    rows = retrieve(user["id"], knowledge_base_id, "knowledge learning", 1)
    source = rows[0]["content"] if rows else "学习需要持续练习并及时复习。"
    knowledge = (source.splitlines()[0].lstrip("# ") or "基础知识")[:80]
    quiz_id, question_id = str(uuid.uuid4()), str(uuid.uuid4())
    question = {"id": question_id, "prompt": f"以下哪项最符合资料内容？\n{source[:180]}", "options": ["资料中的核心描述", "完全相反的说法", "与主题无关", "无法判断"], "answer": "资料中的核心描述", "knowledge": knowledge}
    with db() as conn:
        conn.execute("INSERT INTO quizzes VALUES (?,?,?,?)", (quiz_id, user["id"], knowledge_base_id, now()))
        conn.execute("INSERT INTO quiz_questions VALUES (?,?,?,?,?,?)", (question_id, quiz_id, question["prompt"], json.dumps(question["options"], ensure_ascii=False), question["answer"], knowledge))
    public_question = {key: value for key, value in question.items() if key != "answer"}
    return {"id": quiz_id, "questions": [public_question]}


@app.post("/api/v1/quizzes/{quiz_id}/submit")
def submit_quiz(quiz_id: str, body: SubmitIn, user=Depends(current_user)) -> dict[str, Any]:
    with db() as conn:
        quiz = conn.execute("SELECT * FROM quizzes WHERE id=? AND user_id=?", (quiz_id, user["id"])).fetchone()
        questions = conn.execute("SELECT * FROM quiz_questions WHERE quiz_id=?", (quiz_id,)).fetchall()
        if not quiz:
            raise HTTPException(404, "Quiz not found")
        score = sum(body.answers.get(q["id"]) == q["answer"] for q in questions) / max(1, len(questions)) * 100
        conn.execute("INSERT INTO quiz_attempts VALUES (?,?,?,?,?)", (str(uuid.uuid4()), quiz_id, user["id"], score, now()))
        for q in questions:
            old = conn.execute("SELECT score FROM mastery WHERE user_id=? AND knowledge=?", (user["id"], q["knowledge"])).fetchone()
            new_score = round((old["score"] * 0.6 + score * 0.4) if old else score, 2)
            conn.execute("INSERT INTO mastery VALUES (?,?,?,?) ON CONFLICT(user_id, knowledge) DO UPDATE SET score=excluded.score, updated_at=excluded.updated_at", (user["id"], q["knowledge"], new_score, now()))
    return {"score": score, "mastery_score": new_score, "correct": score == 100}


@app.get("/api/v1/mastery")
def mastery(user=Depends(current_user)) -> list[dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT knowledge,score,updated_at FROM mastery WHERE user_id=? ORDER BY score", (user["id"],)).fetchall()
    return [dict(row) for row in rows]


@app.get("/api/v1/knowledge-graph")
def knowledge_graph(user=Depends(current_user)) -> dict[str, Any]:
    nodes = mastery(user)
    return {"nodes": [{"id": n["knowledge"], "label": n["knowledge"], "mastery": n["score"]} for n in nodes], "edges": []}


@app.get("/api/v1/stats")
def stats(user=Depends(current_user)) -> dict[str, Any]:
    with db() as conn:
        kb = conn.execute("SELECT COUNT(*) FROM knowledge_bases WHERE user_id=?", (user["id"],)).fetchone()[0]
        docs = conn.execute("SELECT COUNT(*) FROM documents d JOIN knowledge_bases k ON k.id=d.knowledge_base_id WHERE k.user_id=?", (user["id"],)).fetchone()[0]
        tasks = conn.execute("SELECT COUNT(*), COALESCE(SUM(completed),0) FROM study_tasks t JOIN study_plans p ON p.id=t.plan_id WHERE p.user_id=?", (user["id"],)).fetchone()
        quizzes = conn.execute("SELECT COUNT(*), COALESCE(AVG(score),0) FROM quiz_attempts WHERE user_id=?", (user["id"],)).fetchone()
    return {"knowledge_bases": kb, "documents": docs, "tasks_total": tasks[0], "tasks_completed": tasks[1], "task_completion_rate": round(tasks[1] / tasks[0] * 100, 2) if tasks[0] else 0, "quiz_attempts": quizzes[0], "quiz_accuracy": round(quizzes[1], 2)}


def index_repository(user_id: str, url: str, root: Path, knowledge_base_id: str | None = None) -> dict[str, Any]:
    parsed = public_repository_url(url)
    project_name = Path(parsed.path.rstrip("/")).name.removesuffix(".git") or parsed.hostname or "repository"
    analysis, source_documents = analyze_repository(root, project_name)
    kb_id = knowledge_base_id
    if kb_id:
        owned_kb(user_id, kb_id)
    else:
        kb_id = str(uuid.uuid4())
        with db() as conn:
            conn.execute("INSERT INTO knowledge_bases VALUES (?,?,?,?)", (kb_id, user_id, project_name[:80], now()))
    indexed = 0
    vector_chunks: list[tuple[str, str, int]] = []
    with db() as conn:
        for filename, content in source_documents:
            if not content.strip():
                continue
            document_id = str(uuid.uuid4())
            chunks = make_chunks(content)
            conn.execute("INSERT INTO documents VALUES (?,?,?,?,?,?,?)", (document_id, kb_id, filename[:180], "text/plain", content, "indexed", now()))
            for position, chunk in enumerate(chunks):
                chunk_id = str(uuid.uuid4())
                conn.execute("INSERT INTO chunks VALUES (?,?,?,?,?)", (chunk_id, document_id, chunk, position, json.dumps({"filename": filename, "section": position + 1, "repository": url}, ensure_ascii=False)))
                vector_chunks.append((chunk_id, chunk, position, document_id))
            indexed += 1
        import_id = str(uuid.uuid4())
        conn.execute("INSERT INTO repository_imports VALUES (?,?,?,?,?,?)", (import_id, user_id, url, "indexed", json.dumps(analysis, ensure_ascii=False), now()))
    vector_indexed = index_chunks(user_id, kb_id, "repository", vector_chunks) if vector_chunks else True
    return {"id": import_id, "url": url, "status": "indexed", "knowledge_base_id": kb_id, "document_count": indexed, "analysis": analysis, "vector_indexed": vector_indexed}


@app.post("/api/v1/repositories/import", status_code=201)
def import_repository(body: RepositoryIn, user=Depends(current_user)) -> dict[str, Any]:
    parsed = public_repository_url(body.url)
    with tempfile.TemporaryDirectory(prefix="knowflow-repo-") as workspace:
        target = Path(workspace) / "repo"
        root = clone_repository(body.url, target)
        return index_repository(user["id"], body.url, root, body.knowledge_base_id)


@app.post("/api/v1/memories", status_code=201)
def save_memory(body: MemoryIn, user=Depends(current_user)) -> dict[str, Any]:
    return create_memory(user["id"], body)


@app.get("/api/v1/memories")
def list_memories(query: str | None = None, user=Depends(current_user)) -> list[dict[str, Any]]:
    return search_memories(user["id"], query)


@app.get("/api/v1/memories/{memory_id}")
def memory_detail(memory_id: str, user=Depends(current_user)) -> dict[str, Any]:
    with db() as conn:
        row = conn.execute("SELECT * FROM memories WHERE id=? AND user_id=?", (memory_id, user["id"])).fetchone()
    if not row:
        raise HTTPException(404, "Memory not found")
    result = dict(row)
    result["metadata"] = json.loads(result["metadata"])
    return result


@app.delete("/api/v1/memories/{memory_id}")
def delete_memory(memory_id: str, user=Depends(current_user)) -> dict[str, bool]:
    with db() as conn:
        cursor = conn.execute("DELETE FROM memories WHERE id=? AND user_id=?", (memory_id, user["id"]))
        if cursor.rowcount == 0:
            raise HTTPException(404, "Memory not found")
    return {"deleted": True}


TOOL_SPECS = {
    "search_knowledge": {"description": "Search the user's indexed knowledge", "write": False, "required": ["query"], "properties": {"query": {"type": "string"}, "knowledge_base_id": {"type": "string"}}},
    "read_document": {"description": "Read an owned indexed document", "write": False, "required": ["document_id"], "properties": {"document_id": {"type": "string"}}},
    "get_learning_stats": {"description": "Read learning statistics", "write": False, "required": [], "properties": {}},
    "get_mastery": {"description": "Read mastery records", "write": False, "required": [], "properties": {}},
    "create_study_plan": {"description": "Create a study plan", "write": True, "required": ["goal"], "properties": {"goal": {"type": "string"}, "days": {"type": "integer", "minimum": 1, "maximum": 14}}},
    "list_tasks": {"description": "List owned learning tasks", "write": False, "required": [], "properties": {}},
    "update_task": {"description": "Update an owned task", "write": True, "required": ["task_id", "completed"], "properties": {"task_id": {"type": "string"}, "completed": {"type": "boolean"}}},
    "generate_quiz": {"description": "Generate a quiz from an owned knowledge base", "write": True, "required": [], "properties": {"knowledge_base_id": {"type": "string"}}},
    "get_knowledge_graph": {"description": "Read the mastery knowledge graph", "write": False, "required": [], "properties": {}},
    "save_memory": {"description": "Save a durable user memory", "write": True, "required": ["content"], "properties": {"memory_type": {"type": "string"}, "content": {"type": "string"}, "importance": {"type": "number"}, "source": {"type": "string"}}},
    "search_memory": {"description": "Search durable user memories", "write": False, "required": [], "properties": {"query": {"type": "string"}}},
    "import_repository": {"description": "Clone and statically analyze a public repository", "write": True, "required": ["url"], "properties": {"url": {"type": "string"}, "knowledge_base_id": {"type": "string"}}},
}


def dispatch_tool(user_id: str, name: str, arguments: dict[str, Any]) -> Any:
    spec = TOOL_SPECS.get(name)
    if not spec:
        raise ValueError("Unknown tool")
    missing = [key for key in spec["required"] if key not in arguments]
    if missing:
        raise ValueError(f"Missing required arguments: {', '.join(missing)}")
    with db() as conn:
        user = conn.execute("SELECT * FROM users WHERE id=?", (user_id,)).fetchone()
    if name == "search_knowledge":
        return {"citations": [citation(row) for row in retrieve(user_id, arguments.get("knowledge_base_id"), arguments["query"])]}
    if name == "read_document":
        return document_detail(arguments["document_id"], user)
    if name == "get_learning_stats":
        return stats(user)
    if name == "get_mastery":
        return mastery(user)
    if name == "create_study_plan":
        return create_plan(PlanIn(goal=arguments["goal"], days=arguments.get("days", 7)), user)
    if name == "list_tasks":
        return list_tasks(user)
    if name == "update_task":
        return update_task(arguments["task_id"], {"completed": arguments["completed"]}, user)
    if name == "generate_quiz":
        return create_quiz(arguments.get("knowledge_base_id"), user)
    if name == "get_knowledge_graph":
        return knowledge_graph(user)
    if name == "save_memory":
        return create_memory(user_id, MemoryIn(**arguments))
    if name == "search_memory":
        return search_memories(user_id, arguments.get("query"))
    if name == "import_repository":
        return import_repository(RepositoryIn(**arguments), user)
    raise ValueError("Unsupported tool")


def provider_configured() -> bool:
    return all(os.getenv(key) for key in ("OPENAI_BASE_URL", "OPENAI_API_KEY", "OPENAI_MODEL"))


def provider_completion(
    messages: list[dict[str, Any]],
    *,
    use_tools: bool = True,
    timeout_seconds: float = 20.0,
) -> dict[str, Any]:
    endpoint = os.getenv("OPENAI_BASE_URL", "").rstrip("/")
    if not endpoint.endswith("/chat/completions"):
        endpoint += "/chat/completions"
    request_body: dict[str, Any] = {"model": os.environ["OPENAI_MODEL"], "messages": messages, "temperature": 0.2}
    if use_tools:
        tools = [{"type": "function", "function": {"name": name, "description": spec["description"], "parameters": {"type": "object", "properties": spec["properties"], "required": spec["required"]}}} for name, spec in TOOL_SPECS.items()]
        request_body["tools"] = tools
        request_body["tool_choice"] = "auto"
    payload = json.dumps(request_body).encode()
    request = Request(endpoint, data=payload, headers={"Authorization": f"Bearer {os.environ['OPENAI_API_KEY']}", "Content-Type": "application/json"}, method="POST")
    try:
        with urlopen(request, timeout=timeout_seconds) as response:
            return json.loads(response.read().decode("utf-8"))
    except Exception as exc:
        raise HTTPException(502, "Configured model provider failed") from exc


CHAT_INSUFFICIENT_EVIDENCE = "NOT_ENOUGH_EVIDENCE"
CHAT_NO_EVIDENCE_ANSWER = "我还没有找到相关资料。请先上传文档，或换一个更具体的问题。"
CHAT_DETERMINISTIC_PREFIX = "基于你的知识库，"
CHAT_TIMEOUT_SECONDS = float(os.getenv("CHAT_TIMEOUT_SECONDS", "60"))
CHAT_SYSTEM_PROMPT = (
    "你是 KnowFlow 学习助手。只能依据用户消息中给出的知识库片段回答，不得使用片段之外的任何知识，"
    "不得编造来源或引用。如果这些片段不足以回答，只回复 NOT_ENOUGH_EVIDENCE，不要输出任何其它内容。"
    "回答要简洁，并使用与提问相同的语言。"
)


class InsufficientEvidenceError(RuntimeError):
    """The configured model judged the retrieved passages insufficient to answer the question."""


def chat_generation_enabled() -> bool:
    """Model-generated answers are opt-in; demo mode keeps the deterministic cited excerpt."""
    return provider_configured() and os.getenv("DEMO_AI_MODE", "1") != "1"


def provider_chat_answer(question: str, rows: list[sqlite3.Row]) -> str:
    """Ask the configured provider for an answer grounded strictly in the retrieved passages."""
    evidence = "\n\n".join(
        f"[{index}] {row['filename']} 第{row['position'] + 1}节：{row['content']}"
        for index, row in enumerate(rows, start=1)
    )
    messages = [
        {"role": "system", "content": CHAT_SYSTEM_PROMPT},
        {"role": "user", "content": f"知识库片段：\n{evidence}\n\n问题：{question}"},
    ]
    response = provider_completion(messages, use_tools=False, timeout_seconds=CHAT_TIMEOUT_SECONDS)
    answer = (response.get("choices", [{}])[0].get("message", {}).get("content") or "").strip()
    if not answer or CHAT_INSUFFICIENT_EVIDENCE in answer:
        raise InsufficientEvidenceError("model reported insufficient evidence")
    return answer


def audit_tool_call(run_id: str, user_id: str, name: str, arguments: dict[str, Any]) -> dict[str, Any]:
    status = "completed"
    try:
        result = dispatch_tool(user_id, name, arguments)
    except (HTTPException, ValueError, TypeError) as exc:
        status, result = "failed", {"error": str(exc.detail if isinstance(exc, HTTPException) else exc)}
    with db() as conn:
        conn.execute("INSERT INTO tool_calls VALUES (?,?,?,?,?,?,?,?)", (str(uuid.uuid4()), run_id, user_id, name, json.dumps(arguments, ensure_ascii=False), status, json.dumps(result, ensure_ascii=False, default=str), now()))
    return {"name": name, "status": status, "result": result}


def provider_agent(user_id: str, run_id: str, body: AgentRunIn) -> tuple[str, list[dict[str, Any]]]:
    messages: list[dict[str, Any]] = [{"role": "system", "content": "Use the supplied tools for KnowFlow user data. Never invent tool results."}, {"role": "user", "content": body.message}]
    results: list[dict[str, Any]] = []
    for _ in range(4):
        response = provider_completion(messages)
        assistant = response.get("choices", [{}])[0].get("message", {})
        calls = assistant.get("tool_calls") or []
        if not calls:
            return assistant.get("content", ""), results
        messages.append(assistant)
        for call in calls:
            function = call.get("function", {})
            name = function.get("name", "")
            try:
                arguments = json.loads(function.get("arguments", "{}"))
            except json.JSONDecodeError:
                arguments = {}
            result = audit_tool_call(run_id, user_id, name, arguments)
            results.append(result)
            messages.append({"role": "tool", "tool_call_id": call.get("id", str(uuid.uuid4())), "content": json.dumps(result["result"], ensure_ascii=False, default=str)})
    raise HTTPException(502, "Model provider exceeded tool-call limit")


@app.get("/api/v1/agents/tools")
def agent_tools(user=Depends(current_user)) -> list[dict[str, Any]]:
    return [{"name": name, "description": spec["description"], "write": spec["write"], "json_schema": {"type": "object", "properties": spec["properties"], "required": spec["required"]}} for name, spec in TOOL_SPECS.items()]


@app.post("/api/v1/agents/run")
def run_agent(body: AgentRunIn, user=Depends(current_user)) -> dict[str, Any]:
    if not body.tool_calls and not provider_configured():
        raise HTTPException(503, "No tool calls supplied and no configured model provider")
    run_id = str(uuid.uuid4())
    with db() as conn:
        conn.execute("INSERT INTO agent_runs VALUES (?,?,?,?,?,?)", (run_id, user["id"], body.conversation_id, body.message, "running", now()))
    if body.tool_calls:
        results = [audit_tool_call(run_id, user["id"], call.name, call.arguments) for call in body.tool_calls]
        content = "工具执行完成" if all(item["status"] == "completed" for item in results) else "部分工具执行失败"
    else:
        content, results = provider_agent(user["id"], run_id, body)
    failed = any(item["status"] != "completed" for item in results)
    with db() as conn:
        conn.execute("UPDATE agent_runs SET status=? WHERE id=?", ("failed" if failed else "completed", run_id))
    return {"run_id": run_id, "content": content, "tool_calls": results}


@app.get("/api/v1/agents/runs")
def agent_runs(user=Depends(current_user)) -> list[dict[str, Any]]:
    with db() as conn:
        rows = conn.execute("SELECT tc.*, ar.message FROM tool_calls tc JOIN agent_runs ar ON ar.id=tc.run_id WHERE tc.user_id=? ORDER BY tc.created_at DESC", (user["id"],)).fetchall()
    return [dict(row) for row in rows]


if __name__ == "__main__":
    import uvicorn
    uvicorn.run("app.main:app", host="0.0.0.0", port=int(os.getenv("PORT", "8001")), reload=False)
